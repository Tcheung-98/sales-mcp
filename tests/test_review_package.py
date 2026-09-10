"""B2 review package builder (PI-2522 / docs/DECK-QA-ARCHITECTURE.md §5).

Offline: the FortuneAI fixture stands in for the S3 template, product clones come
from a stub Presentation, and render_slides is mocked. The LibreOffice render is
a separate, skipped-by-default test.
"""

from __future__ import annotations

import json
import os
import shutil
import zipfile
from datetime import date
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from pptx import Presentation

from ingestion.generator import DeckGenerator
from ingestion.gtm_product_map import GtmProductMap, ProductSlideRef
from ingestion.manifest import manifest_validation_errors
from ingestion.placeholder_fills import apply_placeholders
from ingestion.review_package import (
    PITCH_START_INDEX,
    RENDER_DPI,
    build_review_package,
    plan_pitch_sequence,
)
from ingestion.schema import DeckSchema, Product
from tests.fortuneai_placeholder_fixture import (
    MINIMAL_PNG,
    fortuneai_fixture_bytes,
    mock_placeholder_ai,
    sample_audience_data,
)

_FORTUNEAI_URL = "https://fortune.sharepoint.com/sites/x/FortuneAI_DeckTemplate.pptx"

# Product Name → (GTM Product Category, Deck Path, Slide #). Names and categories
# are taken from tests/gtm_fixtures.py::REPRESENTATIVE_GTM_ROWS.
_GTM_ROWS: dict[str, tuple[str, str, int]] = {
    "CEO Daily": ("Newsletters", "Fortune_Newsletters_2026.pptx", 3),
    "Term Sheet": ("Newsletters", "Fortune_Newsletters_2026.pptx", 7),
    "Crown Unit": ("Digital Ads/Programmatic", "Fortune_Digital_2026.pptx", 5),
    "Long-Form Article": ("Branded Content", "Fortune_BrandedContent_2026.pptx", 2),
    "Full Page": ("Print", "Fortune_Print_2026.pptx", 4),
}


def _gtm_map(*names: str) -> GtmProductMap:
    return GtmProductMap(
        [
            ProductSlideRef(
                product_name=name,
                category=_GTM_ROWS[name][0],
                deck_path=_GTM_ROWS[name][1],
                slide_number=_GTM_ROWS[name][2],
            )
            for name in names
        ]
    )


def _schema(*products: Product, **overrides) -> DeckSchema:
    defaults = dict(
        company_name="Acme Corp",
        industry="Technology",
        budgets=[{"amount": sum(p.price for p in products)}],
        flight_dates={"start": "2026-09-28", "end": "2026-12-31"},
        campaign_goal="Drive consideration among enterprise buyers",
        targeting_details=(
            "US enterprise tech decision-makers, Chief Executive Officer, C-suite"
        ),
        kpis=["Awareness", "Engagement"],
        kpi_details="Lift brand awareness 10%; engagement rate above benchmark",
        campaign_narrative="Acme helps mid-market CFOs modernize finance ops",
        preferred_platforms_products=["Newsletters", "Branded Content"],
        additional_rfp_details="Prefer Q4 flight; avoid holiday blackout weeks",
        client_logo="https://example.com/acme-logo.png",
        confirmed_products=list(products),
    )
    defaults.update(overrides)
    return DeckSchema(**defaults)


def _newsletter_schema() -> DeckSchema:
    return _schema(
        Product(name="CEO Daily", cadence="weekly", price=50_000, category="Newsletter")
    )


def _product_source_prs(n: int = 8) -> Presentation:
    prs = Presentation()
    for i in range(n):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = f"PRODUCT CLONE {i + 1}"
    return prs


def _post_c2_deck(schema: DeckSchema, gtm_map: GtmProductMap) -> Presentation:
    """C1 assembly + C2 fills on the CI fixture — the deck B2 packages (8 + P slides)."""
    with patch("boto3.client", return_value=MagicMock()):
        generator = DeckGenerator(bucket="test-bucket")
    generator._s3 = MagicMock()
    with (
        patch("requests.get") as mock_get,
        patch.object(generator, "_load_pptx", return_value=_product_source_prs()),
    ):
        mock_get.return_value.content = fortuneai_fixture_bytes()
        mock_get.return_value.raise_for_status = MagicMock()
        prs = generator.assemble_skeleton(
            schema, _FORTUNEAI_URL, product_map=gtm_map
        )
    apply_placeholders(
        prs,
        schema,
        audience=sample_audience_data(),
        logo_bytes=MINIMAL_PNG,
        as_of=date(2026, 9, 1),
        ai=mock_placeholder_ai(),
    )
    return prs


def _fake_render(pptx, slide_indices, *, output_dir, dpi, **kwargs):
    """Stand-in for render_slides: same slide-NNN.png naming, no LibreOffice."""
    assert dpi == RENDER_DPI
    assert slide_indices is None  # every slide
    out = Path(output_dir)
    count = len(Presentation(str(pptx)).slides)
    paths = []
    for i in range(count):
        png = out / f"slide-{i:03d}.png"
        png.write_bytes(MINIMAL_PNG)
        paths.append(png)
    return paths


def _build(schema, gtm_map, tmp_path, **kwargs):
    prs = _post_c2_deck(schema, gtm_map)
    with patch("ingestion.review_package.render_slides", side_effect=_fake_render):
        return build_review_package(
            prs, schema, gtm_map=gtm_map, output_dir=tmp_path, **kwargs
        )


def _soffice_available() -> bool:
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    soffice = (
        os.environ.get("SOFFICE_BIN")
        or shutil.which("soffice")
        or shutil.which("libreoffice")
        or (str(mac) if mac.is_file() else None)
    )
    return bool(soffice) and bool(shutil.which("pdftoppm"))


def test_plan_pitch_sequence_orders_dividers_by_workflow():
    schema = _schema(
        Product(name="Full Page", cadence="annual", price=10_000, category="Print"),
        Product(name="Term Sheet", cadence="weekly", price=20_000, category="Newsletter"),
        Product(
            name="Crown Unit", cadence="monthly", price=30_000, category="Digital Media"
        ),
        Product(name="CEO Daily", cadence="weekly", price=40_000, category="Newsletter"),
    )
    plan = plan_pitch_sequence(
        schema, _gtm_map("Full Page", "Term Sheet", "Crown Unit", "CEO Daily")
    )

    # High-Impact (0) → Editorial (1) → Print (3); Premium Video and Branded unfunded.
    assert [kind for kind, _ in plan] == [
        "divider",
        "product",
        "divider",
        "product",
        "product",
        "divider",
        "product",
    ]
    assert [item[1] for item in plan if item[0] == "divider"] == [0, 1, 3]
    assert [item[1].product_name for item in plan if item[0] == "product"] == [
        "Crown Unit",
        "Term Sheet",
        "CEO Daily",
        "Full Page",
    ]


def test_plan_pitch_sequence_keeps_raw_gtm_deck_path():
    plan = plan_pitch_sequence(_newsletter_schema(), _gtm_map("CEO Daily"))
    ref = plan[1][1]
    assert ref.deck_path == "Fortune_Newsletters_2026.pptx"
    assert ref.slide_number == 3


def test_plan_pitch_sequence_fails_loud_on_missing_gtm_row():
    with pytest.raises(ValueError, match="No GTM Product Tags row"):
        plan_pitch_sequence(_newsletter_schema(), GtmProductMap([]))


def test_plan_pitch_sequence_escalates_events():
    schema = _schema(
        Product(
            name="BrainStorm Summit", cadence="annual", price=50_000, category="Events"
        )
    )
    with pytest.raises(ValueError, match="escalate"):
        plan_pitch_sequence(schema, GtmProductMap([]))


def test_single_newsletter_package_matches_architecture_index_map(tmp_path):
    schema = _newsletter_schema()
    package = _build(schema, _gtm_map("CEO Daily"), tmp_path)
    manifest = package.manifest

    # P = 1 divider + 1 product → 10 slides, product clone at index 7 (§5).
    assert manifest.slide_count == 10
    assert [(s.role, s.slide_kind) for s in manifest.slides] == [
        ("cover", None),
        ("narrative", None),
        ("narrative", None),
        ("narrative", None),
        ("narrative", None),
        ("narrative", None),
        ("other", "divider"),
        ("product", None),
        ("other", "investment"),
        ("other", "thank_you"),
    ]
    clone = manifest.slides[7]
    assert clone.editable is False
    assert clone.product_name == "CEO Daily"
    assert clone.source_path == "Fortune_Newsletters_2026.pptx"
    assert clone.source_slide_number == 3
    assert all(s.editable for s in manifest.slides if s.role != "product")


def test_package_writes_manifest_and_schema(tmp_path):
    schema = _newsletter_schema()
    package = _build(schema, _gtm_map("CEO Daily"), tmp_path)

    assert package.root == tmp_path
    assert package.draft_path == tmp_path / "draft.pptx"
    assert package.manifest_path == tmp_path / "manifest.json"
    assert package.slides_dir == tmp_path / "slides"
    assert package.pptx_bytes == package.draft_path.read_bytes()
    assert len(Presentation(str(package.draft_path)).slides) == 10

    data = json.loads(package.manifest_path.read_text())
    assert manifest_validation_errors(data) == []
    assert data["client_name"] == "Acme Corp"
    assert data["template_key"] == "FortuneAI_DeckTemplate.pptx"
    assert data["slides"][7]["editable"] is False

    reloaded = DeckSchema.model_validate_json(package.schema_path.read_text())
    assert reloaded.confirmed_products[0].name == "CEO Daily"


def test_package_renders_one_png_per_slide(tmp_path):
    package = _build(_newsletter_schema(), _gtm_map("CEO Daily"), tmp_path)
    assert [p.name for p in package.slide_pngs] == [
        f"slide-{i:03d}.png" for i in range(10)
    ]
    assert sorted(p.name for p in package.slides_dir.glob("*.png")) == [
        f"slide-{i:03d}.png" for i in range(10)
    ]


def test_multi_category_package_marks_every_clone_uneditable(tmp_path):
    schema = _schema(
        Product(
            name="Crown Unit", cadence="monthly", price=100_000, category="Digital Media"
        ),
        Product(name="CEO Daily", cadence="weekly", price=75_000, category="Newsletter"),
        Product(name="Term Sheet", cadence="weekly", price=60_000, category="Newsletter"),
        Product(
            name="Long-Form Article",
            cadence="quarterly",
            price=95_000,
            category="Branded Content",
        ),
        Product(name="Full Page", cadence="annual", price=70_000, category="Print"),
    )
    gtm_map = _gtm_map(*_GTM_ROWS)
    package = _build(schema, gtm_map, tmp_path)
    manifest = package.manifest

    plan = plan_pitch_sequence(schema, gtm_map)
    assert manifest.slide_count == 8 + len(plan)  # 4 dividers + 5 products
    assert manifest.slide_count == 17

    product_indices = [s.slide_index for s in manifest.slides if s.role == "product"]
    assert product_indices == [
        PITCH_START_INDEX + i for i, (kind, _) in enumerate(plan) if kind == "product"
    ]
    assert all(
        s.editable is False for s in manifest.slides if s.role == "product"
    )
    assert [s.product_name for s in manifest.slides if s.role == "product"] == [
        "Crown Unit",
        "CEO Daily",
        "Term Sheet",
        "Full Page",
        "Long-Form Article",
    ]
    assert [s.slide_kind for s in manifest.slides[-2:]] == ["investment", "thank_you"]


def test_draft_pptx_has_no_duplicate_slide_parts(tmp_path):
    """Clones carry their source partnames; an un-renumbered save collides."""
    schema = _schema(
        Product(
            name="Crown Unit", cadence="monthly", price=100_000, category="Digital Media"
        ),
        Product(name="CEO Daily", cadence="weekly", price=75_000, category="Newsletter"),
        Product(name="Term Sheet", cadence="weekly", price=60_000, category="Newsletter"),
        Product(
            name="Long-Form Article",
            cadence="quarterly",
            price=95_000,
            category="Branded Content",
        ),
        Product(name="Full Page", cadence="annual", price=70_000, category="Print"),
    )
    package = _build(schema, _gtm_map(*_GTM_ROWS), tmp_path)

    names = zipfile.ZipFile(package.draft_path).namelist()
    assert len(names) == len(set(names))
    assert len(Presentation(str(package.draft_path)).slides) == package.manifest.slide_count


def test_build_review_package_requires_plan_or_gtm_map(tmp_path):
    schema = _newsletter_schema()
    prs = _post_c2_deck(schema, _gtm_map("CEO Daily"))
    with pytest.raises(ValueError, match="plan= or gtm_map="):
        build_review_package(prs, schema, output_dir=tmp_path)


def test_build_review_package_rejects_slide_count_mismatch(tmp_path):
    schema = _newsletter_schema()
    gtm_map = _gtm_map("CEO Daily")
    prs = _post_c2_deck(schema, gtm_map)
    plan = plan_pitch_sequence(schema, gtm_map)
    with pytest.raises(ValueError, match="post-C2 deck must have 9 slides"):
        build_review_package(prs, schema, plan=plan[:1], output_dir=tmp_path)


def test_cli_help_exits_clean():
    from ingestion.review_package import _main

    with pytest.raises(SystemExit) as exc:
        _main(["--help"])
    assert exc.value.code == 0


@pytest.mark.skipif(
    not _soffice_available(), reason="needs LibreOffice (SOFFICE_BIN) and pdftoppm"
)
def test_package_renders_real_pngs(tmp_path):
    schema = _newsletter_schema()
    gtm_map = _gtm_map("CEO Daily")
    package = build_review_package(
        _post_c2_deck(schema, gtm_map), schema, gtm_map=gtm_map, output_dir=tmp_path
    )
    assert len(package.slide_pngs) == package.manifest.slide_count
    assert all(p.is_file() and p.stat().st_size > 0 for p in package.slide_pngs)
