"""BambooHR Tier 1 golden fixture (Deck QA architecture §12, option 1).

Product names are verbatim from ``tests.logic_guide_fixtures``
``REPRESENTATIVE_GTM_ROWS`` — the only Product Tags rows the repo confirms
exist — so ``GtmProductMap.lookup`` resolves here and against the live workbook.
Prices are whole multiples of the matching ``REPRESENTATIVE_PRICING_ROWS`` rate
and sum to the stated $400,000 tier; cadences are what
``confirm_mix.derive_cadence`` returns for those same rates.

Assertion-free on purpose: ``tests/test_bamboohr_golden.py`` owns the checks and
other Deck QA tests can reuse the schema.
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.presentation import Presentation as PresentationType

from ingestion.category_dividers import (
    CATEGORY_DIVIDERS,
    FORTUNEAI_DIVIDER_SLIDE_INDEX,
)
from ingestion.gtm_product_map import (
    GtmProductMap,
    ProductSlideRef,
    normalize_category,
    product_deck_s3_key,
)
from ingestion.placeholder_fills import funded_divider_buckets
from ingestion.schema import DeckSchema, Product
from tests.fortuneai_placeholder_fixture import build_fortuneai_fixture_prs
from tests.logic_guide_fixtures import REPRESENTATIVE_GTM_ROWS

BAMBOOHR_TIER1_BUDGET = 400_000.0

# Segments proven to exist in both sample_audience_data() and the live
# Audience Data tab (smoke_build_live.py uses the same three).
BAMBOOHR_TARGETING = (
    "US enterprise HR, people operations, and finance decision-makers: "
    "Chief Executive Officer, Chief Financial Officer, C-suite"
)

# Rate card units behind each price (REPRESENTATIVE_PRICING_ROWS):
#   Crown Unit        $25,000    x 4  = $100,000
#   CEO Daily         $5,000/day x 15 = $75,000
#   Term Sheet        $6,000/day x 10 = $60,000
#   Full Page         $35,000    x 3  = $105,000
#   Long-Form Article $60,000    x 1  = $60,000
BAMBOOHR_PRODUCTS: tuple[Product, ...] = (
    Product(name="Crown Unit", cadence="quarterly", price=100_000, category="Digital Media"),
    Product(name="CEO Daily", cadence="weekly", price=75_000, category="Newsletter"),
    Product(name="Term Sheet", cadence="weekly", price=60_000, category="Newsletter"),
    Product(name="Full Page", cadence="quarterly", price=105_000, category="Print"),
    Product(
        name="Long-Form Article",
        cadence="quarterly",
        price=60_000,
        category="Branded Content",
    ),
)

_PRODUCT_PAGE_SUFFIX = " PRODUCT PAGE"


def bamboohr_tier1_schema() -> DeckSchema:
    """Locked Tier 1 handoff schema for BambooHR ($400,000 mix)."""
    return DeckSchema(
        company_name="BambooHR",
        industry="Technology",
        budgets=[{"amount": BAMBOOHR_TIER1_BUDGET, "label": "Tier 1"}],
        flight_dates={"start": "2026-09-28", "end": "2026-12-31"},
        campaign_goal="Own the HR technology conversation with enterprise decision-makers",
        targeting_details=BAMBOOHR_TARGETING,
        kpis=["Awareness", "Brand Lift"],
        kpi_details="Lift unaided awareness among C-suite buyers; beat brand-lift benchmark",
        campaign_narrative=(
            "BambooHR helps growing companies run people operations without enterprise overhead"
        ),
        preferred_platforms_products=[
            "Digital Ads/Programmatic",
            "Newsletters",
            "Print",
            "Branded Content",
        ],
        additional_rfp_details="Q4 flight; avoid holiday blackout weeks",
        client_logo="https://example.com/bamboohr-logo.png",
        confirmed_products=[product.model_copy() for product in BAMBOOHR_PRODUCTS],
    )


def bamboohr_product_map() -> GtmProductMap:
    """Product Tags rows for the golden mix, copied from REPRESENTATIVE_GTM_ROWS."""
    wanted = {
        (normalize_category(product.category), product.name)
        for product in BAMBOOHR_PRODUCTS
    }
    rows = [
        ProductSlideRef(
            product_name=name,
            category=category,
            deck_path=deck_path,
            slide_number=int(slide_number),
        )
        for category, name, _tags, deck_path, slide_number in REPRESENTATIVE_GTM_ROWS
        if (category, name) in wanted
    ]
    if len(rows) != len(BAMBOOHR_PRODUCTS):
        found = sorted(row.product_name for row in rows)
        raise ValueError(
            "REPRESENTATIVE_GTM_ROWS no longer covers the BambooHR golden mix "
            f"(matched {found})"
        )
    return GtmProductMap(rows)


def expected_pitch_sequence() -> list[tuple[str, str]]:
    """Post-C1 pitch section as ``("divider" | "product", name)`` in Workflow order."""
    sequence: list[tuple[str, str]] = []
    for divider_name, products in funded_divider_buckets(bamboohr_tier1_schema()):
        sequence.append(("divider", divider_name))
        sequence.extend(("product", product.name) for product in products)
    return sequence


def expected_pitch_slide_count() -> int:
    """``P`` from architecture §5: funded dividers + confirmed products."""
    return len(expected_pitch_sequence())


def expected_slide_count() -> int:
    """Final post-C2 deck size: ``8 + P`` (architecture §5)."""
    return 8 + expected_pitch_slide_count()


def product_page_title(product_name: str) -> str:
    """Title carried by the stand-in A5 clone for ``product_name``."""
    return f"{product_name.upper()}{_PRODUCT_PAGE_SUFFIX}"


def expected_pitch_titles() -> list[str]:
    """Slide titles the pitch section should carry, in order."""
    return [
        name if kind == "divider" else product_page_title(name)
        for kind, name in expected_pitch_sequence()
    ]


def bamboohr_template_bytes() -> bytes:
    """CI FortuneAI fixture with dividers re-titled at their real physical indices.

    ``build_fortuneai_fixture_prs`` adds the five dividers in Workflow order, but
    the real template stores them at ``FORTUNEAI_DIVIDER_SLIDE_INDEX``. Re-titling
    by physical index is what makes post-assembly divider order assertable.
    """
    prs = build_fortuneai_fixture_prs()
    for category_index, slide_index in enumerate(FORTUNEAI_DIVIDER_SLIDE_INDEX):
        prs.slides[slide_index].shapes.title.text = CATEGORY_DIVIDERS[category_index].name
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _product_page_titles_by_key() -> dict[str, dict[int, str]]:
    gtm_map = bamboohr_product_map()
    by_key: dict[str, dict[int, str]] = {}
    for product in BAMBOOHR_PRODUCTS:
        ref = gtm_map.lookup(product.name, product.category)
        key = product_deck_s3_key(ref.deck_path)
        by_key.setdefault(key, {})[ref.slide_number] = product_page_title(product.name)
    return by_key


def fake_product_decks() -> dict[str, PresentationType]:
    """Stand-in Hunter product decks keyed by S3 key (no live ``product-decks/``).

    Each deck is padded to its highest referenced Slide # so the A5 clone lands
    on the exact slide the Product Tags row names.
    """
    decks: dict[str, PresentationType] = {}
    for key, titles in _product_page_titles_by_key().items():
        prs = Presentation()
        for slide_number in range(1, max(titles) + 1):
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = titles.get(slide_number, f"UNRELATED SLIDE {slide_number}")
        decks[key] = prs
    return decks


def fake_load_pptx():
    """``side_effect`` for ``DeckGenerator._load_pptx`` covering the golden mix."""
    decks = fake_product_decks()

    def _load(s3_key: str) -> PresentationType:
        try:
            return decks[s3_key]
        except KeyError as exc:
            raise AssertionError(f"unexpected product deck load: {s3_key}") from exc

    return _load
