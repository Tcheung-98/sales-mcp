"""B4 headless Cursor QA runner (offline; fake cursor_sdk, no network)."""

from __future__ import annotations

import json
import sys
import time
import types
from datetime import date
from pathlib import Path

import pytest
from pptx import Presentation

from ingestion.category_dividers import FORTUNEAI_TEMPLATE_BASENAME
from ingestion.deck_qa_agent import (
    DEFAULT_TIMEOUT_S,
    MAX_ATTACHED_IMAGES,
    UNREPORTED_FIX_LABEL,
    load_review_package,
    main,
    resolve_timeout_s,
    run_headless_cursor_qa,
    select_slide_images,
)
from ingestion.manifest import ReviewManifest, SlideManifestEntry
from ingestion.placeholder_fills import apply_placeholders
from ingestion.schema import DeckSchema, Product
from tests.fortuneai_placeholder_fixture import (
    MINIMAL_PNG,
    build_fortuneai_fixture_prs,
    mock_placeholder_ai,
    sample_audience_data,
)

# §5 shape for P = 2: cover, 5 narrative, divider, product clone, investment, thanks.
_TEN_SLIDE_ROLES: list[tuple[str, str | None]] = [
    ("cover", None),
    *[("narrative", None)] * 5,
    ("other", "divider"),
    ("product", None),
    ("other", "investment"),
    ("other", "thank_you"),
]


def _schema(**overrides) -> DeckSchema:
    defaults = dict(
        company_name="Acme Corp",
        industry="Technology",
        budgets=[{"amount": 50_000}],
        flight_dates={"start": "2026-09-01", "end": "2026-12-31"},
        campaign_goal="Drive consideration among enterprise buyers",
        targeting_details="Chief Executive Officer, C-suite, Chief Financial Officer",
        kpis=["Awareness", "Engagement"],
        kpi_details="Lift brand awareness 10%; engagement rate above benchmark",
        campaign_narrative="Acme helps mid-market CFOs modernize finance ops",
        preferred_platforms_products=["Newsletters"],
        additional_rfp_details="Prefer Q4 flight",
        client_logo="https://example.com/acme-logo.png",
        confirmed_products=[
            Product(
                name="CEO Daily", cadence="weekly", price=50_000, category="Newsletter"
            )
        ],
    )
    defaults.update(overrides)
    return DeckSchema(**defaults)


def _stub_deck(slide_count: int) -> Presentation:
    prs = Presentation()
    for index in range(slide_count):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = f"STUB SLIDE {index}"
    return prs


def _manifest_entries(roles: list[tuple[str, str | None]]) -> list[SlideManifestEntry]:
    return [
        SlideManifestEntry(slide_index=index, role=role, slide_kind=kind)
        for index, (role, kind) in enumerate(roles)
    ]


def _write_package(
    root: Path,
    *,
    roles: list[tuple[str, str | None]] | None = None,
    prs: Presentation | None = None,
    manifest_slide_count: int | None = None,
    schema: DeckSchema | None = None,
) -> Path:
    roles = roles if roles is not None else _TEN_SLIDE_ROLES
    schema = schema if schema is not None else _schema()
    deck = prs if prs is not None else _stub_deck(len(roles))
    root.mkdir(parents=True, exist_ok=True)
    deck.save(str(root / "draft.pptx"))

    manifest = ReviewManifest(
        client_name=schema.client_name,
        template_key=FORTUNEAI_TEMPLATE_BASENAME,
        slide_count=manifest_slide_count or len(roles),
        slides=_manifest_entries(roles),
    )
    (root / "manifest.json").write_text(
        manifest.model_dump_json(indent=2, exclude_none=True), encoding="utf-8"
    )
    (root / "deck_schema.json").write_text(
        schema.model_dump_json(indent=2), encoding="utf-8"
    )
    slides = root / "slides"
    slides.mkdir(exist_ok=True)
    for index in range(len(roles)):
        (slides / f"slide-{index:03d}.png").write_bytes(MINIMAL_PNG)
    return root


def _fortuneai_package(root: Path) -> tuple[Path, DeckSchema]:
    schema = _schema()
    prs = build_fortuneai_fixture_prs()
    apply_placeholders(
        prs,
        schema,
        audience=sample_audience_data(),
        logo_bytes=MINIMAL_PNG,
        as_of=date(2026, 8, 18),
        ai=mock_placeholder_ai(),
    )
    roles: list[tuple[str, str | None]] = [
        ("cover" if i == 0 else "narrative" if i <= 5 else "other", None)
        for i in range(len(prs.slides))
    ]
    return _write_package(root, roles=roles, prs=prs, schema=schema), schema


def _install_fake_sdk(monkeypatch, *, on_run=None, status: str = "finished"):
    log = types.SimpleNamespace(
        creates=0, prompts=[], images=[], options=[], cancels=0, closes=0
    )

    class SDKImage:
        def __init__(self, path: str) -> None:
            self.path = path

        @classmethod
        def from_file(cls, path, *, mime_type=None, dimension=None) -> "SDKImage":
            return cls(str(path))

    class UserMessage:
        def __init__(self, text: str, images=None) -> None:
            self.text = text
            self.images = list(images or [])

    class LocalAgentOptions:
        def __init__(self, cwd=None, **_kwargs) -> None:
            self.cwd = cwd

    class AgentOptions:
        def __init__(self, model=None, api_key=None, local=None, **_kwargs) -> None:
            self.model = model
            self.api_key = api_key
            self.local = local

    class Run:
        id = "run-test"

        def wait(self):
            if on_run is not None:
                on_run()
            return types.SimpleNamespace(status=status)

        def supports(self, _operation: str) -> bool:
            return True

        def cancel(self) -> None:
            log.cancels += 1

    class Agent:
        agent_id = "agent-test"

        @classmethod
        def create(cls, options=None, **_kwargs) -> "Agent":
            log.creates += 1
            log.options.append(options)
            return cls()

        def __enter__(self) -> "Agent":
            return self

        def __exit__(self, *_exc) -> bool:
            log.closes += 1
            return False

        def send(self, message, options=None) -> Run:
            log.prompts.append(message.text)
            log.images.extend(image.path for image in message.images)
            return Run()

    module = types.ModuleType("cursor_sdk")
    module.Agent = Agent
    module.AgentOptions = AgentOptions
    module.LocalAgentOptions = LocalAgentOptions
    module.SDKImage = SDKImage
    module.UserMessage = UserMessage
    monkeypatch.setitem(sys.modules, "cursor_sdk", module)
    monkeypatch.setenv("CURSOR_API_KEY", "test-key")
    return log


def _agent_writes(root: Path, payload: dict, *, touch_draft: bool = False):
    def _run() -> None:
        if touch_draft:
            draft = root / "draft.pptx"
            draft.write_bytes(draft.read_bytes() + b"\x00")
        (root / "qa_cursor.json").write_text(json.dumps(payload), encoding="utf-8")

    return _run


def _attached_indices(log) -> list[int]:
    return [int(Path(path).stem.split("-")[-1]) for path in log.images]


def test_deterministic_failure_skips_the_agent(tmp_path, monkeypatch):
    log = _install_fake_sdk(monkeypatch)
    root = _write_package(tmp_path / "pkg", manifest_slide_count=11)

    assert main(["--package-dir", str(root)]) == 1
    assert log.creates == 0
    assert not (root / "qa_cursor.json").exists()

    det = json.loads((root / "qa_deterministic.json").read_text())
    assert det["passed"] is False
    assert "slide_count" in [c["name"] for c in det["checks"] if not c["passed"]]


def test_force_runs_the_agent_and_records_the_override(tmp_path, monkeypatch):
    root = _write_package(tmp_path / "pkg", manifest_slide_count=11)
    log = _install_fake_sdk(
        monkeypatch,
        on_run=_agent_writes(
            root, {"passed": True, "loop_count": 0, "issues": [], "fixes_applied": []}
        ),
    )

    assert main(["--package-dir", str(root), "--force"]) == 0
    assert log.creates == 1

    report = json.loads((root / "qa_cursor.json").read_text())
    assert report["issues"][0]["severity"] == "warning"
    assert "--force" in report["issues"][0]["message"]


def test_cli_success_path_runs_both_gates(tmp_path, monkeypatch):
    root, _ = _fortuneai_package(tmp_path / "pkg")
    log = _install_fake_sdk(
        monkeypatch,
        on_run=_agent_writes(
            root, {"passed": True, "loop_count": 0, "issues": [], "fixes_applied": []}
        ),
    )

    assert main(["--package-dir", str(root)]) == 0
    assert log.creates == 1
    assert json.loads((root / "qa_deterministic.json").read_text())["passed"] is True
    assert json.loads((root / "qa_cursor.json").read_text())["passed"] is True


@pytest.mark.parametrize(
    ("touch_draft", "agent_payload", "expected_fixes", "expected_loop"),
    [
        (
            True,
            {"passed": True, "loop_count": 0, "issues": [], "fixes_applied": []},
            [UNREPORTED_FIX_LABEL],
            1,
        ),
        (
            False,
            {
                "passed": True,
                "loop_count": 1,
                "issues": [],
                "fixes_applied": ["intro_title"],
            },
            [],
            1,
        ),
        (
            True,
            {
                "passed": True,
                "loop_count": 1,
                "issues": [],
                "fixes_applied": ["opportunity_body"],
            },
            ["opportunity_body"],
            1,
        ),
    ],
)
def test_fix_reconciliation(
    tmp_path,
    monkeypatch,
    touch_draft,
    agent_payload,
    expected_fixes,
    expected_loop,
):
    root = _write_package(tmp_path / "pkg")
    package = load_review_package(root)
    _install_fake_sdk(
        monkeypatch,
        on_run=_agent_writes(root, agent_payload, touch_draft=touch_draft),
    )

    report = run_headless_cursor_qa(package, timeout_s=30)

    assert report.fixes_applied == expected_fixes
    assert report.loop_count == expected_loop


def test_product_slide_pngs_are_never_attached(tmp_path, monkeypatch):
    root = _write_package(tmp_path / "pkg")
    package = load_review_package(root)
    log = _install_fake_sdk(
        monkeypatch,
        on_run=_agent_writes(
            root, {"passed": True, "loop_count": 0, "issues": [], "fixes_applied": []}
        ),
    )

    run_headless_cursor_qa(package, timeout_s=30)

    assert _attached_indices(log) == [0, 1, 2, 3, 4, 5, 6, 8, 9]
    assert all("slide-007.png" not in path for path in log.images)


def test_image_attachments_are_capped(tmp_path):
    roles = [("cover", None), *[("narrative", None)] * 5]
    roles += [("other", "divider"), ("product", None)] * 6
    roles += [("other", "investment"), ("other", "thank_you")]
    root = _write_package(tmp_path / "pkg", roles=roles)
    package = load_review_package(root)

    images = select_slide_images(package)
    indices = [int(png.stem.split("-")[-1]) for png in images]

    assert len(images) == MAX_ATTACHED_IMAGES
    assert indices == sorted(indices)
    assert indices[:6] == [0, 1, 2, 3, 4, 5]
    assert len(roles) - 2 in indices and len(roles) - 1 in indices
    product_indices = {e.slide_index for e in package.manifest.slides if not e.editable}
    assert product_indices.isdisjoint(indices)


def test_product_clone_error_is_downgraded(tmp_path, monkeypatch):
    root = _write_package(tmp_path / "pkg")
    package = load_review_package(root)
    _install_fake_sdk(
        monkeypatch,
        on_run=_agent_writes(
            root,
            {
                "passed": True,
                "loop_count": 0,
                "issues": [
                    {
                        "slide_index": 7,
                        "severity": "error",
                        "message": "CEO Daily clone has a typo in its subhead.",
                    }
                ],
                "fixes_applied": [],
            },
        ),
    )

    report = run_headless_cursor_qa(package, timeout_s=30)

    assert report.passed is True
    assert [(i.slide_index, i.severity) for i in report.issues] == [(7, "warning")]


def test_error_on_editable_slide_overrides_a_claimed_pass(tmp_path, monkeypatch):
    root = _write_package(tmp_path / "pkg")
    package = load_review_package(root)
    _install_fake_sdk(
        monkeypatch,
        on_run=_agent_writes(
            root,
            {
                "passed": True,
                "loop_count": 1,
                "issues": [
                    {
                        "slide_index": 3,
                        "severity": "error",
                        "message": "[BODY] still rendered on the Opportunity slide.",
                    }
                ],
                "fixes_applied": [],
            },
        ),
    )

    report = run_headless_cursor_qa(package, timeout_s=30)

    assert report.passed is False
    assert "failed" in report.summary()


@pytest.mark.parametrize(
    "failure_mode",
    ["missing_report", "sdk_error", "failed_status"],
)
def test_agent_failure_modes(tmp_path, monkeypatch, failure_mode):
    root = _write_package(tmp_path / "pkg")
    package = load_review_package(root)

    if failure_mode == "missing_report":
        _install_fake_sdk(monkeypatch)
    elif failure_mode == "sdk_error":
        _install_fake_sdk(monkeypatch)
        module = sys.modules["cursor_sdk"]

        def _boom_create(cls, options=None, **_kwargs):
            raise RuntimeError("sdk unavailable")

        monkeypatch.setattr(module.Agent, "create", classmethod(_boom_create))
    else:
        _install_fake_sdk(
            monkeypatch,
            status="error",
            on_run=_agent_writes(
                root,
                {"passed": True, "loop_count": 0, "issues": [], "fixes_applied": []},
            ),
        )

    report = run_headless_cursor_qa(package, timeout_s=30)

    assert report.passed is False
    assert report.issues


def test_timeout_cancels_the_run_and_fails(tmp_path, monkeypatch):
    root = _write_package(tmp_path / "pkg")
    package = load_review_package(root)
    log = _install_fake_sdk(monkeypatch, on_run=lambda: time.sleep(0.3))

    report = run_headless_cursor_qa(package, timeout_s=0.05)

    assert report.passed is False
    assert log.cancels == 1
    assert any("timed out" in issue.message for issue in report.issues)


def test_missing_api_key_fails_loud(tmp_path, monkeypatch):
    root = _write_package(tmp_path / "pkg")
    package = load_review_package(root)
    _install_fake_sdk(monkeypatch)
    monkeypatch.delenv("CURSOR_API_KEY", raising=False)

    with pytest.raises(ValueError, match="CURSOR_API_KEY"):
        run_headless_cursor_qa(package, timeout_s=30)


def test_prompt_carries_the_skill_and_the_manifest(tmp_path, monkeypatch):
    root = _write_package(tmp_path / "pkg")
    package = load_review_package(root)
    log = _install_fake_sdk(
        monkeypatch,
        on_run=_agent_writes(
            root, {"passed": True, "loop_count": 0, "issues": [], "fixes_applied": []}
        ),
    )

    run_headless_cursor_qa(package, timeout_s=30)
    prompt = log.prompts[0]

    assert "B4 headless vision pass" in prompt
    assert "Never edit a slide with `editable: false`" in prompt
    assert str(package.draft_path) in prompt
    assert "role=product" in prompt
    assert "Flag-only slides, never edit them: [7]" in prompt
    assert log.options[0].local.cwd == str(Path(__file__).resolve().parents[1])


def test_load_review_package_requires_the_b2_layout(tmp_path):
    with pytest.raises(ValueError, match="draft.pptx"):
        load_review_package(tmp_path)


def test_resolve_timeout_prefers_argument_then_env(monkeypatch):
    monkeypatch.delenv("DECK_QA_TIMEOUT_S", raising=False)
    assert resolve_timeout_s() == DEFAULT_TIMEOUT_S
    monkeypatch.setenv("DECK_QA_TIMEOUT_S", "42")
    assert resolve_timeout_s() == 42
    assert resolve_timeout_s(7) == 7
    monkeypatch.setenv("DECK_QA_TIMEOUT_S", "soon")
    with pytest.raises(ValueError, match="DECK_QA_TIMEOUT_S"):
        resolve_timeout_s()
