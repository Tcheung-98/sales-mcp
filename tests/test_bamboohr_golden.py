"""BambooHR golden deck test (Deck QA architecture §5 geometry, §12 fixture).

The assertion this golden exists for is ``slide_count == 8 + P`` with the §5
index map: dividers in Workflow order, each A5 clone under its own divider,
investment and thank-you as the last two slides. The dollar figures are not the
point.

Offline by default — no S3, no AWS credentials, no Anthropic, no Cursor SDK::

    PYTHONPATH=. uv run pytest tests/test_bamboohr_golden.py -q

The two ``live`` tests hit the real FortuneAI template, GTM workbook and cached
product decks (Claude still mocked). They skip unless ``S3_SNAPSHOT_BUCKET`` is
set and those assets are readable::

    PYTHONPATH=. S3_SNAPSHOT_BUCKET=<bucket> uv run pytest \
        tests/test_bamboohr_golden.py -q -k live
"""

from __future__ import annotations

import io
import os
from datetime import date
from unittest.mock import MagicMock, patch

import boto3
import pytest
from pptx import Presentation

from ingestion.generator import DECK_QA_BYPASSED_WARNING, DeckGenerator
from ingestion.gtm_product_map import load_gtm_product_map_from_s3
from ingestion.placeholder_fills import (
    AUDIENCE_SEGMENT_TOKEN,
    AUDIENCE_TITLE_TOKEN,
    BODY_TOKEN,
    BUDGET_TOKEN,
    CLIENT_NAME_UPPER_TOKEN,
    DATE_TOKEN,
    EM_DASH,
    HEADER_TOKEN,
    INDEX_TOKEN,
    PRODUCT_CATEGORY_TOKEN,
    PRODUCT_DESCRIPTION_LITERAL,
    PRODUCT_TYPE_LITERAL,
    REACH_TOKEN,
    TITLE_TOKEN,
    apply_placeholders,
    format_usd,
)
from ingestion.pptx_tools import (
    CLIENT_NAME_POSSESSIVE_TOKEN,
    CLIENT_NAME_TOKEN,
    LOGO_TOKEN,
    iter_shapes,
)
from tests.bamboohr_golden import (
    BAMBOOHR_PRODUCTS,
    BAMBOOHR_TIER1_BUDGET,
    bamboohr_product_map,
    bamboohr_template_bytes,
    bamboohr_tier1_schema,
    expected_pitch_slide_count,
    expected_pitch_titles,
    expected_slide_count,
    fake_load_pptx,
)
from tests.fortuneai_placeholder_fixture import (
    MINIMAL_PNG,
    mock_placeholder_ai,
    sample_audience_data,
)

_FORTUNEAI_URL = "https://fortune.sharepoint.com/sites/x/FortuneAI_DeckTemplate.pptx"
GOLDEN_AS_OF = date(2026, 9, 8)

# Architecture §6 leftover-token list, from the source constants.
LEFTOVER_TOKENS = (
    TITLE_TOKEN,
    HEADER_TOKEN,
    BODY_TOKEN,
    AUDIENCE_TITLE_TOKEN,
    DATE_TOKEN,
    AUDIENCE_SEGMENT_TOKEN,
    REACH_TOKEN,
    INDEX_TOKEN,
    BUDGET_TOKEN,
    PRODUCT_CATEGORY_TOKEN,
    PRODUCT_TYPE_LITERAL,
    PRODUCT_DESCRIPTION_LITERAL,
    CLIENT_NAME_UPPER_TOKEN,
    CLIENT_NAME_TOKEN,
    CLIENT_NAME_POSSESSIVE_TOKEN,
    LOGO_TOKEN,
    # Investment slide; PR-B adds the constant (architecture §6).
    "[PRODUCT + PRODUCT PRICE]",
)


def _slide_text(slide) -> str:
    parts: list[str] = []
    for shape in iter_shapes(slide):
        if not getattr(shape, "has_text_frame", False):
            continue
        for para in shape.text_frame.paragraphs:
            parts.append(para.text or "")
    return "\n".join(parts)


def _deck_text(prs) -> str:
    return "\n".join(_slide_text(slide) for slide in prs.slides)


def _slide_titles(prs) -> list[str]:
    return [
        slide.shapes.title.text if slide.shapes.title is not None else ""
        for slide in prs.slides
    ]


def _generator() -> DeckGenerator:
    with patch("boto3.client", return_value=MagicMock()):
        generator = DeckGenerator(bucket="test-bucket")
    generator._s3 = MagicMock()
    generator._api_key = "test-key"
    return generator


def _assemble_golden(generator: DeckGenerator, schema):
    with (
        patch("requests.get") as mock_get,
        patch.object(generator, "_load_pptx", side_effect=fake_load_pptx()),
    ):
        mock_get.return_value.content = bamboohr_template_bytes()
        mock_get.return_value.raise_for_status = MagicMock()
        return generator.assemble_skeleton(
            schema, _FORTUNEAI_URL, product_map=bamboohr_product_map()
        )


def _golden_deck(schema):
    """Assembled (C1) then mock-filled (C2) golden deck; returns (prs, warnings)."""
    prs = _assemble_golden(_generator(), schema)
    warnings = apply_placeholders(
        prs,
        schema,
        audience=sample_audience_data(),
        logo_bytes=MINIMAL_PNG,
        as_of=GOLDEN_AS_OF,
        ai=mock_placeholder_ai(),
    )
    return prs, warnings


def test_assembled_skeleton_keeps_workflow_pitch_order():
    """Pre-C2 skeleton is 14 + P with the pitch section starting at index 12."""
    pitch = expected_pitch_slide_count()
    prs = _assemble_golden(_generator(), bamboohr_tier1_schema())

    titles = _slide_titles(prs)
    assert len(prs.slides) == 14 + pitch
    assert titles[12 : 12 + pitch] == expected_pitch_titles()
    assert "Premium Video" not in titles  # unfunded divider omitted


def test_golden_deck_slide_count_and_index_map():
    """Post-C2 geometry: 8 + P, pitch at 6 … 5 + P, investment and thanks last."""
    pitch = expected_pitch_slide_count()
    prs, warnings = _golden_deck(bamboohr_tier1_schema())

    assert warnings == []
    assert len(prs.slides) == expected_slide_count()
    assert len(prs.slides) == 8 + pitch
    assert _slide_titles(prs)[6 : 6 + pitch] == expected_pitch_titles()
    assert format_usd(BAMBOOHR_TIER1_BUDGET) in _slide_text(prs.slides[-2])
    assert "Thank you!" in _slide_text(prs.slides[-1])


def test_golden_deck_passes_basic_checks():
    """Mock-filled deck: no leftover tokens, every product named, re-opens clean."""
    prs, _warnings = _golden_deck(bamboohr_tier1_schema())
    blob = _deck_text(prs)

    for token in LEFTOVER_TOKENS:
        assert token not in blob, f"leftover token {token!r}"
    investment = _slide_text(prs.slides[-2])
    for product in BAMBOOHR_PRODUCTS:
        assert product.name.lower() in blob.lower()
        assert f"{product.name} {EM_DASH} {format_usd(product.price)}" in investment
    assert "BambooHR" in blob
    assert "September 2026" in blob

    DeckGenerator._renumber_slide_parts(prs)  # build() does this before saving
    buf = io.BytesIO()
    prs.save(buf)
    assert len(Presentation(io.BytesIO(buf.getvalue())).slides) == expected_slide_count()


def test_build_returns_golden_slide_count(monkeypatch):
    """The slide_count build_deck hands back is 8 + P."""
    # Geometry test only — bypass the always-on QA rail (no LibreOffice/Cursor here).
    monkeypatch.setenv("DECK_QA_DISABLED", "1")
    generator = _generator()
    generator._s3.generate_presigned_url.return_value = "https://s3.example.com/deck.pptx"

    with (
        patch("requests.get") as mock_get,
        patch.object(generator, "_load_pptx", side_effect=fake_load_pptx()),
        patch(
            "ingestion.generator.PlaceholderAI.from_anthropic",
            return_value=mock_placeholder_ai(),
        ),
    ):
        mock_get.return_value.content = bamboohr_template_bytes()
        mock_get.return_value.raise_for_status = MagicMock()
        result = generator.build(
            bamboohr_tier1_schema(),
            _FORTUNEAI_URL,
            product_map=bamboohr_product_map(),
            audience_data=sample_audience_data(),
            logo_bytes=MINIMAL_PNG,
        )

    assert result["slide_count"] == expected_slide_count()
    assert result["client_name"] == "BambooHR"
    assert result["warnings"] == [DECK_QA_BYPASSED_WARNING]


def _skip_if_no_bucket() -> str:
    bucket = os.environ.get("S3_SNAPSHOT_BUCKET")
    if not bucket:
        pytest.skip("S3_SNAPSHOT_BUCKET not set")
    return bucket


def test_live_golden_products_resolve_in_gtm_workbook():
    """Every golden product name is a real Product Tags row (§12 verification)."""
    bucket = _skip_if_no_bucket()
    try:
        gtm_map = load_gtm_product_map_from_s3(boto3.client("s3"), bucket)
    except Exception as exc:
        pytest.skip(f"S3 GTM workbook unavailable: {exc}")

    for product in BAMBOOHR_PRODUCTS:
        ref = gtm_map.lookup(product.name, product.category)
        assert ref.deck_path
        assert ref.slide_number >= 1


def test_live_golden_deck_matches_slide_geometry():
    """Real template + real A5 clones still land on 8 + P (Claude mocked)."""
    bucket = _skip_if_no_bucket()
    schema = bamboohr_tier1_schema()
    try:
        generator = DeckGenerator(bucket=bucket)
        gtm_map = generator._get_gtm_product_map()
        audience = generator._get_audience_data()
    except Exception as exc:
        pytest.skip(f"S3 GTM workbook unavailable: {exc}")

    try:
        prs = generator.assemble_skeleton(schema, template_url=None, product_map=gtm_map)
    except ValueError as exc:
        if "Failed to load" in str(exc):
            pytest.skip(f"S3 deck asset unavailable: {exc}")
        raise

    apply_placeholders(
        prs,
        schema,
        audience=audience,
        logo_bytes=MINIMAL_PNG,
        as_of=GOLDEN_AS_OF,
        ai=mock_placeholder_ai(),
    )

    assert len(prs.slides) == expected_slide_count()
    blob = _deck_text(prs)
    for token in LEFTOVER_TOKENS:
        assert token not in blob, f"leftover token {token!r}"
