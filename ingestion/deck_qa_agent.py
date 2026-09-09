"""B4 headless Cursor QA (docs/DECK-QA-ARCHITECTURE.md §7).

Drives a Cursor SDK agent over a B2 review package: the deck-qa skill becomes the
agent's instructions, the rendered PNGs of **editable** slides are attached as
images, and the agent gets one fix pass (Hard Rule 7) before writing its verdict.

Lives in ``ingestion/`` rather than ``scripts/`` on purpose. ``[tool.setuptools]
packages`` ships ``ingestion`` only, so ``scripts/run_deck_qa.py`` is not
importable from an installed wheel or the Docker image; PR-F calls
:func:`run_headless_cursor_qa` in-process from ``ingestion/generator.py`` (§8) and
``scripts/run_deck_qa.py`` is a thin CLI shim over :func:`main`.

The agent's own ``qa_cursor.json`` is treated as a claim, not as truth. This
module re-reads it, applies the §7 field semantics (severity enum, product-clone
issues capped at ``warning`` and never failing the deck, ``loop_count`` ≤ 1) and
reconciles ``fixes_applied`` against the draft's digest on disk before rewriting
the file. §8 keys its deck re-load off that field: fixes written but not reported
would be silently thrown away and the unfixed deck shipped.
"""

from __future__ import annotations

import hashlib
import json
import logging
import os
from concurrent.futures import ThreadPoolExecutor
from concurrent.futures import TimeoutError as FutureTimeoutError
from dataclasses import dataclass, field
from pathlib import Path

from ingestion.manifest import ReviewManifest, SlideManifestEntry
from ingestion.review_package import (
    CURSOR_REPORT_NAME,
    DECK_SCHEMA_NAME,
    DETERMINISTIC_REPORT_NAME,
    DRAFT_NAME,
    MANIFEST_NAME,
    SLIDES_DIRNAME,
    ReviewPackage,
)

logger = logging.getLogger(__name__)

TIMEOUT_ENV = "DECK_QA_TIMEOUT_S"
API_KEY_ENV = "CURSOR_API_KEY"
DEFAULT_TIMEOUT_S = 600.0
DEFAULT_MODEL = "composer-2.5"

# The agent's instruction set (PR-C). Never duplicated here — read from disk.
SKILL_RELPATH = Path(".cursor") / "skills" / "deck-qa" / "SKILL.md"
REPO_ROOT = Path(__file__).resolve().parents[1]

# Vision budget. Every PNG stays on disk, so the agent can open the rest itself.
MAX_ATTACHED_IMAGES = 12
MAX_FIX_LOOPS = 1

SEVERITIES = ("info", "warning", "error")
_DEFAULT_SEVERITY = "warning"
# Label used when the agent edited draft.pptx but reported no fix (§7).
UNREPORTED_FIX_LABEL = "unreported_draft_edit"
# Grace period for run.wait() to return after a timeout cancel.
_CANCEL_GRACE_S = 30.0

_TAIL_SLIDE_KINDS = ("investment", "thank_you")


@dataclass(frozen=True)
class QaIssue:
    """One ``qa_cursor.json`` issue (§7)."""

    slide_index: int | None
    severity: str
    message: str

    def to_json(self) -> dict:
        return {
            "slide_index": self.slide_index,
            "severity": self.severity,
            "message": self.message,
        }


@dataclass
class CursorQaReport:
    """``qa_cursor.json`` payload (§7).

    Distinct from ``ingestion.deck_qa.QaReport``, which carries B3's named checks.
    Exposes the same ``passed`` / ``summary()`` / ``to_json()`` surface so §8's
    integration code can treat either report the same way.
    """

    passed: bool = False
    loop_count: int = 0
    issues: list[QaIssue] = field(default_factory=list)
    fixes_applied: list[str] = field(default_factory=list)

    @property
    def errors(self) -> list[QaIssue]:
        return [issue for issue in self.issues if issue.severity == "error"]

    def to_json(self) -> dict:
        return {
            "passed": self.passed,
            "loop_count": self.loop_count,
            "issues": [issue.to_json() for issue in self.issues],
            "fixes_applied": list(self.fixes_applied),
        }

    def summary(self) -> str:
        fixes = ", ".join(self.fixes_applied) or "none"
        if self.passed:
            return (
                f"Headless Cursor deck QA passed (loop_count={self.loop_count}, "
                f"{len(self.issues)} issue(s), fixes: {fixes})"
            )
        detail = "; ".join(
            f"slide {issue.slide_index}: {issue.message}"
            if issue.slide_index is not None
            else issue.message
            for issue in self.errors
        )
        return "Headless Cursor deck QA failed: " + (detail or "no error detail reported")


def resolve_timeout_s(timeout_s: float | None = None) -> float:
    """Explicit argument, else ``DECK_QA_TIMEOUT_S``, else 600 (§8)."""
    if timeout_s is not None:
        return float(timeout_s)
    raw = os.environ.get(TIMEOUT_ENV, "").strip()
    if not raw:
        return DEFAULT_TIMEOUT_S
    try:
        return float(raw)
    except ValueError as exc:
        raise ValueError(f"{TIMEOUT_ENV} must be a number, got {raw!r}") from exc


def load_review_package(package_dir: Path | str) -> ReviewPackage:
    """Rebuild a :class:`ReviewPackage` from a package directory already on disk."""
    root = Path(package_dir)
    draft_path = root / DRAFT_NAME
    manifest_path = root / MANIFEST_NAME
    slides_dir = root / SLIDES_DIRNAME
    for path in (draft_path, manifest_path):
        if not path.is_file():
            raise ValueError(f"review package {root} is missing {path.name}")
    if not slides_dir.is_dir():
        raise ValueError(f"review package {root} is missing {SLIDES_DIRNAME}/")

    manifest = ReviewManifest.model_validate_json(
        manifest_path.read_text(encoding="utf-8")
    )
    return ReviewPackage(
        root=root,
        draft_path=draft_path,
        manifest_path=manifest_path,
        slides_dir=slides_dir,
        schema_path=root / DECK_SCHEMA_NAME,
        manifest=manifest,
        pptx_bytes=draft_path.read_bytes(),
        slide_pngs=tuple(sorted(slides_dir.glob("slide-*.png"))),
    )


def run_headless_cursor_qa(
    package: ReviewPackage,
    *,
    timeout_s: float | None = None,
    model: str = DEFAULT_MODEL,
    api_key: str | None = None,
    repo_root: Path | str = REPO_ROOT,
    forced: bool = False,
) -> CursorQaReport:
    """Run the B4 vision pass over ``package`` and return its normalised report.

    ``forced`` records that the caller ran B4 despite a deterministic failure
    (``--force``, dev only). Raises before the agent starts on a missing skill
    file or missing credentials; a run that starts and then fails, times out or
    skips its report comes back as ``passed: false`` so the caller can attach it.
    """
    budget = resolve_timeout_s(timeout_s)
    root = Path(repo_root)
    skill_path = root / SKILL_RELPATH
    if not skill_path.is_file():
        raise ValueError(f"deck-qa skill not found at {skill_path}")
    key = api_key or os.environ.get(API_KEY_ENV, "").strip()
    if not key:
        raise ValueError(f"{API_KEY_ENV} is required to run headless deck QA (§8)")

    images = select_slide_images(package)
    prompt = _build_prompt(
        package,
        skill=skill_path.read_text(encoding="utf-8"),
        images=images,
        repo_root=root,
        forced=forced,
    )

    report_path = package.root / CURSOR_REPORT_NAME
    report_path.unlink(missing_ok=True)
    before = _digest(package.draft_path)

    try:
        status, timed_out = _run_agent_session(
            prompt,
            images,
            cwd=root,
            model=model,
            api_key=key,
            timeout_s=budget,
        )
    except Exception as exc:
        logger.exception("B4 agent session failed")
        report = CursorQaReport(
            passed=False,
            issues=[
                QaIssue(None, "error", f"B4 agent session failed: {exc}"),
            ],
        )
        package.write_cursor_report(report)
        return report

    report = _read_agent_report(report_path, package.manifest)
    if timed_out:
        report.passed = False
        report.issues.append(
            QaIssue(None, "error", f"B4 timed out after {budget:g}s and was cancelled")
        )
    elif status not in ("finished", ""):
        report.passed = False
        report.issues.append(
            QaIssue(None, "error", f"B4 agent run ended with status {status!r}")
        )
    if forced:
        report.issues.insert(
            0,
            QaIssue(
                None,
                "warning",
                "deterministic QA (B3) did not pass; B4 ran under --force",
            ),
        )

    _reconcile_fixes(report, changed=_digest(package.draft_path) != before)
    package.write_cursor_report(report)
    logger.info(
        "B4 report: passed=%s loop_count=%d issues=%d fixes=%s",
        report.passed,
        report.loop_count,
        len(report.issues),
        report.fixes_applied or "none",
    )
    return report


def select_slide_images(package: ReviewPackage) -> list[Path]:
    """PNGs to attach: editable slides only, capped at :data:`MAX_ATTACHED_IMAGES`.

    Product clones are never attached — the agent may not fix them (Hard Rule 1),
    and their PNGs would spend the vision budget on slides it can only flag. It
    still reads them from ``slides/`` when it needs to report on one.
    """
    candidates = [
        (entry, package.slides_dir / f"slide-{entry.slide_index:03d}.png")
        for entry in package.manifest.slides
        if entry.editable
    ]
    missing = [png.name for _, png in candidates if not png.is_file()]
    if missing:
        logger.warning("B4: no rendered PNG for %s", ", ".join(missing))
    present = [(entry, png) for entry, png in candidates if png.is_file()]

    if len(present) > MAX_ATTACHED_IMAGES:
        present.sort(key=lambda pair: (_image_priority(pair[0]), pair[0].slide_index))
        present = present[:MAX_ATTACHED_IMAGES]
        present.sort(key=lambda pair: pair[0].slide_index)
    return [png for _, png in present]


def _image_priority(entry: SlideManifestEntry) -> int:
    """Cover and narrative copy first, then the tail pages, then dividers."""
    if entry.role == "cover":
        return 0
    if entry.role == "narrative":
        return 1
    if entry.slide_kind in _TAIL_SLIDE_KINDS:
        return 2
    return 3


def _digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest() if path.is_file() else ""


def _build_prompt(
    package: ReviewPackage,
    *,
    skill: str,
    images: list[Path],
    repo_root: Path,
    forced: bool,
) -> str:
    manifest = package.manifest
    rows = []
    for entry in manifest.slides:
        row = (
            f"- {entry.slide_index:>3}  role={entry.role}"
            f"  slide_kind={entry.slide_kind or '-'}"
            f"  editable={str(entry.editable).lower()}"
        )
        if entry.product_name:
            row += f"  product_name={entry.product_name!r}"
        rows.append(row)
    locked = [e.slide_index for e in manifest.slides if not e.editable]
    attached = [int(png.stem.split("-")[-1]) for png in images]

    lines = [
        skill.strip(),
        "",
        "---",
        "",
        "# This run",
        "",
        f"Review package: {package.root}",
        f"  deck: {package.draft_path}",
        f"  manifest: {package.manifest_path}",
        f"  schema: {package.schema_path}",
        f"  slide PNGs: {package.slides_dir}",
        f"  B3 report: {package.root / DETERMINISTIC_REPORT_NAME}",
        f"Repo checkout (your cwd): {repo_root}",
        "",
        f"{manifest.client_name} deck, {manifest.slide_count} slides:",
        *rows,
        "",
        f"Attached screenshots (editable slides only): {attached or 'none'}",
        f"Flag-only slides, never edit them: {locked or 'none'}",
        "",
        "Every slide PNG is on disk in that slides/ directory, including the ones "
        "not attached here — read any of them when you need to. Stay inside the "
        "package directory and this checkout.",
        "",
        f"You get {MAX_FIX_LOOPS} fix pass. Re-rendering and re-viewing the slides "
        "you changed is part of it, not a second loop. Edit "
        f"{package.draft_path} in place; do not write a second copy of the deck.",
        "",
        f"Write your verdict to {package.root / CURSOR_REPORT_NAME} in the Exit "
        "contract shape above, then stop. The caller re-reads that file and "
        "reconciles it against the deck's digest on disk: list a slot id in "
        "fixes_applied for every change you saved, and leave it empty if you "
        "saved none.",
    ]
    if forced:
        lines += [
            "",
            "This is a --force dev run: qa_deterministic.json did not pass. Note "
            "that in issues[] as the skill's 'When to use' section requires.",
        ]
    return "\n".join(lines)


def _run_agent_session(
    prompt: str,
    images: list[Path],
    *,
    cwd: Path,
    model: str,
    api_key: str,
    timeout_s: float,
) -> tuple[str, bool]:
    """Send one prompt to a local Cursor agent. Returns ``(status, timed_out)``.

    Imported lazily so the SDK is only needed when QA actually runs, and so the
    test suite can stand a fake ``cursor_sdk`` in its place.
    """
    from cursor_sdk import (
        Agent,
        AgentOptions,
        LocalAgentOptions,
        SDKImage,
        UserMessage,
    )

    message = UserMessage(
        text=prompt, images=[SDKImage.from_file(str(png)) for png in images]
    )
    options = AgentOptions(
        model=model,
        api_key=api_key,
        local=LocalAgentOptions(cwd=str(cwd)),
    )
    with Agent.create(options) as agent:
        run = agent.send(message)
        logger.info(
            "B4 agent %s run %s: %d image(s), %gs budget",
            getattr(agent, "agent_id", "?"),
            getattr(run, "id", "?"),
            len(images),
            timeout_s,
        )
        result = _wait_for_run(run, timeout_s)
        if result is None:
            return "", True
        return str(getattr(result, "status", "") or ""), False


def _wait_for_run(run, timeout_s: float):
    """``run.wait()`` under a wall-clock budget; cancels and returns None on timeout."""
    pool = ThreadPoolExecutor(max_workers=1, thread_name_prefix="deck-qa")
    try:
        future = pool.submit(run.wait)
        try:
            return future.result(timeout=timeout_s)
        except FutureTimeoutError:
            logger.error("B4 agent run exceeded %gs; cancelling", timeout_s)
            if run.supports("cancel"):
                run.cancel()
                # Give the cancelled run a moment to surface its terminal state so
                # the pool thread exits instead of outliving this call.
                try:
                    future.result(timeout=_CANCEL_GRACE_S)
                except Exception:
                    logger.debug("B4 run did not settle after cancel", exc_info=True)
            return None
    finally:
        pool.shutdown(wait=False)


def _read_agent_report(report_path: Path, manifest: ReviewManifest) -> CursorQaReport:
    """Parse the agent's qa_cursor.json and apply the §7 field semantics."""
    if not report_path.is_file():
        return CursorQaReport(
            issues=[
                QaIssue(
                    None,
                    "error",
                    f"B4 agent did not write {report_path.name}",
                )
            ]
        )
    try:
        raw = json.loads(report_path.read_text(encoding="utf-8"))
    except ValueError as exc:
        return CursorQaReport(
            issues=[QaIssue(None, "error", f"{report_path.name} is not valid JSON: {exc}")]
        )
    if not isinstance(raw, dict):
        return CursorQaReport(
            issues=[QaIssue(None, "error", f"{report_path.name} is not a JSON object")]
        )

    editable = {entry.slide_index: entry.editable for entry in manifest.slides}
    issues = [
        _normalize_issue(item, editable) for item in raw.get("issues") or [] if item
    ]
    fixes = [
        str(fix).strip() for fix in raw.get("fixes_applied") or [] if str(fix).strip()
    ]
    loop_count = min(max(_as_int(raw.get("loop_count"), 0), 0), MAX_FIX_LOOPS)
    # A claimed pass cannot survive an error on a slide the agent was allowed to fix.
    blocking = [
        issue
        for issue in issues
        if issue.severity == "error"
        and (issue.slide_index is None or editable.get(issue.slide_index, True))
    ]
    return CursorQaReport(
        passed=bool(raw.get("passed")) and not blocking,
        loop_count=loop_count,
        issues=issues,
        fixes_applied=fixes,
    )


def _normalize_issue(item, editable: dict[int, bool]) -> QaIssue:
    if not isinstance(item, dict):
        return QaIssue(None, _DEFAULT_SEVERITY, str(item))
    slide_index = item.get("slide_index")
    index = _as_int(slide_index, None) if slide_index is not None else None
    severity = str(item.get("severity", "")).strip().lower()
    if severity not in SEVERITIES:
        severity = _DEFAULT_SEVERITY
    # Product clones are flag-only (§7): the pass cannot fix them, so an error
    # there would block delivery on something no one in this pass can repair.
    if severity == "error" and index is not None and not editable.get(index, True):
        severity = "warning"
    return QaIssue(index, severity, str(item.get("message", "")).strip())


def _as_int(value, default):
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def _reconcile_fixes(report: CursorQaReport, *, changed: bool) -> None:
    """Make ``fixes_applied`` non-empty iff draft.pptx changed on disk (§7).

    §8 re-loads the deck only when the list is non-empty, so an unreported edit
    would be discarded and the unfixed deck delivered under a green report.
    """
    if changed and not report.fixes_applied:
        report.fixes_applied = [UNREPORTED_FIX_LABEL]
        report.loop_count = max(report.loop_count, 1)
        report.issues.append(
            QaIssue(
                None,
                "warning",
                "draft.pptx changed on disk but the agent reported no fixes; "
                f"recorded as {UNREPORTED_FIX_LABEL!r} so the deck is re-loaded",
            )
        )
    elif not changed and report.fixes_applied:
        claimed = ", ".join(report.fixes_applied)
        report.fixes_applied = []
        report.issues.append(
            QaIssue(
                None,
                "warning",
                f"agent reported fixes ({claimed}) but draft.pptx is unchanged; "
                "cleared fixes_applied",
            )
        )


def main(argv: list[str] | None = None) -> int:
    """CLI behind ``scripts/run_deck_qa.py``: B3, then B4, over a package directory.

    Exit codes: 0 both passed, 1 deterministic QA failed (B4 skipped), 2 B4 failed.
    """
    import argparse

    from pptx import Presentation

    from ingestion.deck_qa import run_deterministic_qa
    from ingestion.schema import DeckSchema

    parser = argparse.ArgumentParser(
        description=(
            "Run deterministic (B3) then headless Cursor (B4) QA on a review package."
        )
    )
    parser.add_argument(
        "--package-dir",
        type=Path,
        required=True,
        help=f"Review package directory ({DRAFT_NAME}, {MANIFEST_NAME}, {SLIDES_DIRNAME}/)",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="Run B4 even when deterministic QA fails (dev only)",
    )
    parser.add_argument(
        "--timeout-s",
        type=float,
        default=None,
        help=f"Wall-clock budget (default: ${TIMEOUT_ENV}, else {DEFAULT_TIMEOUT_S:g})",
    )
    parser.add_argument("--model", default=DEFAULT_MODEL, help="Cursor SDK model id")
    parser.add_argument(
        "--log-level", default="INFO", help="Root log level (default: INFO)"
    )
    args = parser.parse_args(argv)
    logging.basicConfig(level=args.log_level.upper())

    package = load_review_package(args.package_dir)
    schema = DeckSchema.model_validate_json(
        package.schema_path.read_text(encoding="utf-8")
    )
    det_report = run_deterministic_qa(
        Presentation(str(package.draft_path)), schema, package.manifest
    )
    package.write_deterministic_report(det_report)
    print(det_report.summary())
    if not det_report.passed and not args.force:
        return 1

    report = run_headless_cursor_qa(
        package,
        timeout_s=args.timeout_s,
        model=args.model,
        forced=not det_report.passed,
    )
    print(report.summary())
    for issue in report.issues:
        print(f"  [{issue.severity}] slide {issue.slide_index}: {issue.message}")
    return 0 if report.passed else 2
