"""B3 deterministic deck QA (docs/DECK-QA-ARCHITECTURE.md §6).

Runs on the post-C2 deck, before the headless Cursor vision pass, so a deck
that is already structurally broken never spends a vision loop. Cheap and
offline: no LibreOffice, no S3, no Cursor.

Geometry per §5 — the final deck is ``8 + P`` slides with
``P = funded dividers + confirmed products``: index 0 cover, 1–5 narrative,
6…5+P pitch (dividers + A5 clones), then investment and thank you. The pitch
range is derived from the tail here rather than from ``P`` so the checks also
hold for CI fixtures that were never driven through C1 assembly.

Product clones are scanned but never edited (Hard Rule 1) — B3 only reports.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field

from pptx import Presentation

from ingestion.manifest import ReviewManifest
from ingestion.placeholder_fills import (
    AUDIENCE_SEGMENT_TOKEN,
    AUDIENCE_TITLE_TOKEN,
    BODY_TOKEN,
    BUDGET_TOKEN,
    CLIENT_NAME_UPPER_TOKEN,
    DATE_TOKEN,
    HEADER_TOKEN,
    INDEX_TOKEN,
    PRODUCT_CATEGORY_TOKEN,
    PRODUCT_DESCRIPTION_LITERAL,
    PRODUCT_PRICE_TOKEN,
    PRODUCT_TYPE_LITERAL,
    REACH_TOKEN,
    TITLE_TOKEN,
    format_usd,
    funded_divider_buckets,
    mix_total,
    stated_total_budget,
)
from ingestion.pptx_tools import (
    CLIENT_NAME_POSSESSIVE_TOKEN,
    CLIENT_NAME_TOKEN,
    LOGO_TOKEN,
    iter_shapes,
)
from ingestion.schema import DeckSchema

# First pitch slide (§5 index map); the last two slides are investment + thanks.
PITCH_START_INDEX = 6
TAIL_SLIDE_COUNT = 2
_MIN_SLIDES = PITCH_START_INDEX + TAIL_SLIDE_COUNT
_THANK_YOU_MARKER = "thank you"
# Same tolerance as placeholder_fills._assert_budget_matches_mix.
_BUDGET_TOLERANCE = 0.005
# The summary ends up in a DeckQaError message; keep it readable on a broken deck.
_MAX_REPORTED_HITS = 12

# §6 leftover-token list. Imported, never retyped: CLIENT_NAME_POSSESSIVE_TOKEN
# carries a curly apostrophe and will not match a straight quote.
LEFTOVER_TOKENS: tuple[str, ...] = (
    TITLE_TOKEN,
    HEADER_TOKEN,
    BODY_TOKEN,
    AUDIENCE_TITLE_TOKEN,
    DATE_TOKEN,
    AUDIENCE_SEGMENT_TOKEN,
    REACH_TOKEN,
    INDEX_TOKEN,
    BUDGET_TOKEN,
    PRODUCT_CATEGORY_TOKEN,
    PRODUCT_PRICE_TOKEN,
    PRODUCT_TYPE_LITERAL,
    PRODUCT_DESCRIPTION_LITERAL,
    CLIENT_NAME_UPPER_TOKEN,
    CLIENT_NAME_TOKEN,
    CLIENT_NAME_POSSESSIVE_TOKEN,
    LOGO_TOKEN,
)


@dataclass(frozen=True)
class QaCheckResult:
    """One deterministic check outcome."""

    name: str
    passed: bool
    message: str = ""

    def to_json(self) -> dict:
        return {"name": self.name, "passed": self.passed, "message": self.message}


@dataclass
class QaReport:
    """qa_deterministic.json payload (§5 review-package layout)."""

    checks: list[QaCheckResult] = field(default_factory=list)

    @property
    def passed(self) -> bool:
        return all(check.passed for check in self.checks)

    @property
    def failures(self) -> list[QaCheckResult]:
        return [check for check in self.checks if not check.passed]

    def to_json(self) -> dict:
        return {
            "passed": self.passed,
            "checks": [check.to_json() for check in self.checks],
        }

    def summary(self) -> str:
        if self.passed:
            return f"Deterministic deck QA passed ({len(self.checks)} checks)"
        detail = "; ".join(f"{c.name}: {c.message}" for c in self.failures)
        return f"Deterministic deck QA failed: {detail}"


class DeckQaError(ValueError):
    """QA gate failure carrying its report (§8 failure semantics).

    Subclasses ValueError so server.py's existing ``except ValueError`` arm keeps
    returning ``status: error`` until it learns to attach ``qa_report``.
    """

    def __init__(self, message: str, *, report: QaReport) -> None:
        super().__init__(message)
        self.report = report


def run_deterministic_qa(
    prs,
    schema: DeckSchema,
    manifest: ReviewManifest,
) -> QaReport:
    """Run every §6 check against a post-C2 deck. Never mutates ``prs``."""
    texts = [_slide_text(slide) for slide in prs.slides]
    funded_dividers, divider_error = _funded_divider_names(schema)
    return QaReport(
        checks=[
            _check_slide_count(prs, manifest),
            _check_leftover_tokens(texts),
            _check_investment_budget(schema),
            _check_product_presence(texts, schema),
            _check_divider_order(texts, funded_dividers, divider_error),
            _check_tail_slides(texts, funded_dividers),
            _check_pptx_integrity(prs),
            _check_manifest_consistency(prs, manifest),
        ]
    )


def _slide_text(slide) -> str:
    parts: list[str] = []
    for shape in iter_shapes(slide):
        if not getattr(shape, "has_text_frame", False):
            continue
        for para in shape.text_frame.paragraphs:
            parts.append(para.text or "")
    return "\n".join(parts)


def _funded_divider_names(schema: DeckSchema) -> tuple[list[str], str | None]:
    """Funded divider names in Workflow order, or the reason there are none."""
    try:
        return [name for name, _ in funded_divider_buckets(schema)], None
    except ValueError as exc:
        return [], str(exc)


def _pitch_indices(slide_count: int) -> range:
    """Indices 6 … 5+P, derived from the tail so fixtures work too."""
    return range(PITCH_START_INDEX, max(slide_count - TAIL_SLIDE_COUNT, PITCH_START_INDEX))


def _check_slide_count(prs, manifest: ReviewManifest) -> QaCheckResult:
    actual = len(prs.slides)
    if actual != manifest.slide_count:
        return QaCheckResult(
            "slide_count",
            False,
            f"deck has {actual} slides; manifest declares {manifest.slide_count}",
        )
    return QaCheckResult("slide_count", True, f"{actual} slides")


def _check_leftover_tokens(texts: list[str]) -> QaCheckResult:
    hits = [
        f"slide {index}: {token!r}"
        for index, text in enumerate(texts)
        for token in LEFTOVER_TOKENS
        if token in text
    ]
    if hits:
        shown = hits[:_MAX_REPORTED_HITS]
        overflow = len(hits) - len(shown)
        detail = "; ".join(shown)
        if overflow:
            detail += f"; (+{overflow} more)"
        return QaCheckResult("leftover_tokens", False, f"unfilled placeholders — {detail}")
    return QaCheckResult("leftover_tokens", True, f"{len(LEFTOVER_TOKENS)} tokens clear")


def _check_investment_budget(schema: DeckSchema) -> QaCheckResult:
    """Stated total budget vs mix total.

    C2's ``_assert_budget_matches_mix`` already raises on this inside
    ``apply_placeholders``, so this can never fail on the ``build_deck`` path.
    It bites for standalone package runs (``scripts/run_deck_qa.py
    --package-dir``), where nothing upstream has validated the schema.
    """
    stated = stated_total_budget(schema)
    mix = mix_total(schema)
    if abs(stated - mix) > _BUDGET_TOLERANCE:
        return QaCheckResult(
            "investment_budget",
            False,
            f"stated total budget {format_usd(stated)} does not match mix total "
            f"{format_usd(mix)}",
        )
    return QaCheckResult("investment_budget", True, f"mix totals {format_usd(mix)}")


def _check_product_presence(texts: list[str], schema: DeckSchema) -> QaCheckResult:
    blob = "\n".join(texts).lower()
    missing = [
        product.name
        for product in schema.confirmed_products
        if product.name.lower() not in blob
    ]
    if missing:
        return QaCheckResult(
            "product_presence",
            False,
            f"confirmed products absent from deck text: {', '.join(missing)}",
        )
    return QaCheckResult(
        "product_presence",
        True,
        f"{len(schema.confirmed_products)} confirmed product(s) present",
    )


def _check_divider_order(
    texts: list[str],
    funded_dividers: list[str],
    divider_error: str | None,
) -> QaCheckResult:
    """Funded dividers appear in Workflow order inside the pitch range."""
    if divider_error is not None:
        return QaCheckResult("divider_order", False, divider_error)

    pitch = _pitch_indices(len(texts))
    found: list[tuple[str, int]] = []
    missing: list[str] = []
    for name in funded_dividers:
        index = next(
            (i for i in pitch if name.lower() in texts[i].lower()), None
        )
        if index is None:
            missing.append(name)
        else:
            found.append((name, index))
    if missing:
        return QaCheckResult(
            "divider_order",
            False,
            f"funded divider(s) missing from the pitch section: {', '.join(missing)}",
        )

    indices = [index for _, index in found]
    if indices != sorted(indices):
        order = ", ".join(f"{name}@{index}" for name, index in found)
        return QaCheckResult(
            "divider_order",
            False,
            f"dividers are out of Workflow order: {order}",
        )
    return QaCheckResult(
        "divider_order", True, f"{len(found)} funded divider(s) in Workflow order"
    )


def _check_tail_slides(texts: list[str], funded_dividers: list[str]) -> QaCheckResult:
    """Last two slides are the stock investment + thank-you pages."""
    if len(texts) < _MIN_SLIDES:
        return QaCheckResult(
            "tail_slides",
            False,
            f"deck has {len(texts)} slides; fewer than the {_MIN_SLIDES} stock pages",
        )
    investment, thanks = texts[-2].lower(), texts[-1].lower()
    problems: list[str] = []
    if _THANK_YOU_MARKER not in thanks:
        problems.append("last slide is not the thank-you page")
    if "$" not in investment:
        problems.append("investment slide has no dollar amount")
    absent = [name for name in funded_dividers if name.lower() not in investment]
    if absent:
        problems.append(
            f"investment slide is missing category box(es): {', '.join(absent)}"
        )
    if problems:
        return QaCheckResult("tail_slides", False, "; ".join(problems))
    return QaCheckResult("tail_slides", True, "investment + thank you are last")


def _check_pptx_integrity(prs) -> QaCheckResult:
    """Save to bytes and re-open; catches orphaned rels and section entries."""
    buf = io.BytesIO()
    try:
        prs.save(buf)
        data = buf.getvalue()
        reopened = Presentation(io.BytesIO(data))
    except Exception as exc:
        # Any save or parse failure is a QA failure, not a crash: the report has
        # to reach the caller so the package still records why B4 was skipped.
        return QaCheckResult(
            "pptx_integrity", False, f"deck does not round-trip through python-pptx: {exc}"
        )
    if len(reopened.slides) != len(prs.slides):
        return QaCheckResult(
            "pptx_integrity",
            False,
            f"round-trip lost slides: {len(prs.slides)} saved, "
            f"{len(reopened.slides)} re-opened",
        )
    return QaCheckResult("pptx_integrity", True, f"{len(data)} bytes re-opened cleanly")


def _check_manifest_consistency(prs, manifest: ReviewManifest) -> QaCheckResult:
    slide_count = len(prs.slides)
    out_of_range: list[int] = []
    seen: set[int] = set()
    duplicates: set[int] = set()
    for entry in manifest.slides:
        if entry.slide_index >= slide_count:
            out_of_range.append(entry.slide_index)
        if entry.slide_index in seen:
            duplicates.add(entry.slide_index)
        seen.add(entry.slide_index)

    problems: list[str] = []
    if out_of_range:
        problems.append(
            f"slide_index out of range (deck has {slide_count} slides): "
            f"{sorted(out_of_range)}"
        )
    if duplicates:
        problems.append(f"duplicate slide_index entries: {sorted(duplicates)}")
    if problems:
        return QaCheckResult("manifest_consistency", False, "; ".join(problems))
    return QaCheckResult(
        "manifest_consistency", True, f"{len(manifest.slides)} entries in range"
    )
